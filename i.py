from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

# Criar documento
doc = Document()

# Capa
doc.add_paragraph("\n\n\n")
title = doc.add_paragraph("Levantamento de Requisitos para Sistema Web (EduPlus)")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.runs[0].font.size = Pt(18)
title.runs[0].bold = True

name = doc.add_paragraph("\nVinicius Silva Gonçalves de Almeida")
name.alignment = WD_ALIGN_PARAGRAPH.CENTER
name.runs[0].font.size = Pt(14)

doc.add_page_break()

# Corpo do trabalho
doc.add_heading("🧩 Atividade – Levantamento de Requisitos para Sistema Web (EduPlus)", level=1)
doc.add_paragraph("🎯 Objetivo: Levantar os requisitos funcionais e não funcionais do sistema web de cursos online da empresa EduPlus, pensando nas três personas: aluno, instrutor e administrador.")

# Tarefa 1
doc.add_heading("🧩 Tarefa 1 – Requisitos Funcionais (RF)", level=2)
rf_table = doc.add_table(rows=1, cols=3)
hdr_cells = rf_table.rows[0].cells
hdr_cells[0].text = "Código"
hdr_cells[1].text = "Requisito Funcional"
hdr_cells[2].text = "Tipo"

rfs = [
    ("RF01", "O aluno deve poder se cadastrar e fazer login na plataforma.", "Usuário"),
    ("RF02", "O instrutor deve poder criar novos cursos e adicionar aulas.", "Usuário"),
    ("RF03", "O aluno deve poder se inscrever em cursos disponíveis.", "Usuário"),
    ("RF04", "O aluno deve poder assistir às aulas e fazer avaliações.", "Usuário"),
    ("RF05", "O instrutor deve poder acompanhar o desempenho dos alunos.", "Usuário"),
    ("RF06", "O administrador deve poder gerenciar usuários (alunos e instrutores).", "Sistema"),
    ("RF07", "O administrador deve poder gerar relatórios de desempenho dos cursos.", "Sistema"),
    ("RF08", "O sistema deve enviar notificações (por e-mail ou na plataforma) sobre novas aulas ou avaliações.", "Sistema"),
    ("RF09", "O aluno deve poder acompanhar seu progresso no curso (percentual concluído).", "Usuário"),
    ("RF10", "O sistema deve permitir recuperação de senha por e-mail.", "Sistema")
]

for rf in rfs:
    row_cells = rf_table.add_row().cells
    row_cells[0].text = rf[0]
    row_cells[1].text = rf[1]
    row_cells[2].text = rf[2]

# Tarefa 2
doc.add_heading("⚙ Tarefa 2 – Requisitos Não Funcionais (RNF)", level=2)
rnf_table = doc.add_table(rows=1, cols=4)
hdr_cells = rnf_table.rows[0].cells
hdr_cells[0].text = "Código"
hdr_cells[1].text = "Requisito Não Funcional"
hdr_cells[2].text = "Categoria"
hdr_cells[3].text = "Tipo"

rnfs = [
    ("RNF01", "O sistema deve responder às ações do usuário em até 3 segundos.", "Desempenho", "Produto"),
    ("RNF02", "O sistema deve usar autenticação segura (senha criptografada e token de sessão).", "Segurança", "Produto"),
    ("RNF03", "O sistema deve ser responsivo, funcionando bem em computadores e celulares.", "Portabilidade", "Produto"),
    ("RNF04", "A interface deve ser fácil de usar, com navegação intuitiva.", "Usabilidade", "Produto"),
    ("RNF05", "O sistema deve estar disponível 99% do tempo.", "Confiabilidade", "Produto"),
    ("RNF06", "Somente administradores podem acessar relatórios completos de desempenho.", "Segurança", "Organizacional"),
    ("RNF07", "O sistema deve fazer backup automático diário do banco de dados.", "Confiabilidade", "Organizacional"),
    ("RNF08", "O sistema deve seguir as leis de proteção de dados (LGPD).", "Segurança", "Externo")
]

for rnf in rnfs:
    row_cells = rnf_table.add_row().cells
    row_cells[0].text = rnf[0]
    row_cells[1].text = rnf[1]
    row_cells[2].text = rnf[2]
    row_cells[3].text = rnf[3]

# Tarefa 3
doc.add_heading("🔗 Tarefa 3 – Matriz de Rastreabilidade", level=2)
matrix_table = doc.add_table(rows=1, cols=3)
hdr_cells = matrix_table.rows[0].cells
hdr_cells[0].text = "RF"
hdr_cells[1].text = "RNF Relacionado"
hdr_cells[2].text = "Justificativa"

matrix_data = [
    ("RF01", "RNF02", "O login precisa de segurança para proteger as contas dos usuários."),
    ("RF02", "RNF04", "O instrutor precisa de uma interface simples pra montar as aulas sem complicação."),
    ("RF03", "RNF01", "O sistema deve ser rápido pra não travar quando o aluno se inscreve."),
    ("RF04", "RNF03", "O aluno pode assistir no celular, então o sistema precisa se adaptar bem."),
    ("RF05", "RNF07", "Os dados de desempenho não podem ser perdidos, então precisam de backup."),
    ("RF06", "RNF08", "O administrador precisa seguir as regras de privacidade de dados."),
    ("RF07", "RNF06", "Só o admin pode ver relatórios completos por segurança."),
    ("RF09", "RNF05", "O aluno precisa acessar seu progresso a qualquer hora."),
    ("RF10", "RNF02", "O envio de e-mail de recuperação deve ser protegido.")
]

for m in matrix_data:
    row_cells = matrix_table.add_row().cells
    row_cells[0].text = m[0]
    row_cells[1].text = m[1]
    row_cells[2].text = m[2]

# Tarefa 4
doc.add_heading("⚖ Tarefa 4 – Resolução de Conflitos", level=2)
conflicts_table = doc.add_table(rows=1, cols=3)
hdr_cells = conflicts_table.rows[0].cells
hdr_cells[0].text = "Conflito"
hdr_cells[1].text = "Descrição"
hdr_cells[2].text = "Solução Proposta"

conflicts = [
    ("RF04 x RNF01", "Vídeos em alta qualidade podem deixar o site lento.", "Adicionar opção de escolher a qualidade do vídeo conforme a internet do aluno."),
    ("RF07 x RNF06", "Administradores podem precisar compartilhar dados, mas há limitação de acesso.", "Criar exportação de relatórios com dados anônimos pra preservar a privacidade.")
]

for c in conflicts:
    row_cells = conflicts_table.add_row().cells
    row_cells[0].text = c[0]
    row_cells[1].text = c[1]
    row_cells[2].text = c[2]

# Salvar documento
file_path = "Levantamento_de_Requisitos_Vinicius_Almeida.docx"
doc.save(file_path)

print("Documento gerado com sucesso:", file_path)